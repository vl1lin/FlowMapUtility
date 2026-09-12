import csv

import numpy as np
import pytest
from openpyxl import load_workbook

from flowmaputility.export.grid_exporter import GridExporter
from flowmaputility.grid.info import GridInfo


def _make_grid() -> GridInfo:
    vsl_1d = np.array([0.1, 0.2])
    vsg_1d = np.array([1.0, 2.0])
    vsl_2d = np.array([[0.1, 0.2], [0.1, 0.2]])
    vsg_2d = np.array([[1.0, 1.0], [2.0, 2.0]])
    return GridInfo(vsl_1d, vsg_1d, vsl_2d, vsg_2d, 2, True)


def _make_codes() -> np.ndarray:
    return np.array([[101, 103], [105, 101]])


def test_for_rows():
    grid = _make_grid()
    codes = _make_codes()
    exporter = GridExporter(grid, codes)
    rows = list(exporter._rows())
    assert rows == [
        (1.0, 101, 0.1),
        (1.0, 103, 0.2),
        (2.0, 105, 0.1),
        (2.0, 101, 0.2),
    ]


def test_for_export_to_csv(tmp_path):
    grid = _make_grid()
    codes = _make_codes()
    path = tmp_path / "grid.csv"
    GridExporter(grid, codes).export(str(path))

    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        rows = list(reader)

    assert rows[0] == ["Скорость газа, м/с", "Код режима", "Скорость жидкости, м/с"]
    assert rows[1] == ["1.0", "101", "0.1"]
    assert len(rows) == 5


def test_for_export_to_excel(tmp_path):
    grid = _make_grid()
    codes = _make_codes()
    path = tmp_path / "grid.xlsx"
    GridExporter(grid, codes).export(str(path))

    wb = load_workbook(path)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))  # type: ignore

    assert rows[0] == ("Скорость газа, м/с", "Код режима", "Скорость жидкости, м/с")
    assert rows[1] == (1.0, 101, 0.1)
    assert len(rows) == 5


def test_for_export_unsupported_extension(tmp_path):
    grid = _make_grid()
    codes = _make_codes()
    path = tmp_path / "grid.txt"
    with pytest.raises(ValueError):
        GridExporter(grid, codes).export(str(path))
